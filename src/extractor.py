import re
from pathlib import Path

import fitz
from docx import Document


SKILLS = [
    "Python", "Java", "C++", "JavaScript", "TypeScript",
    "SQL", "MySQL", "PostgreSQL", "MongoDB", "BigQuery",
    "HTML", "CSS", "React", "Node.js", "Flask", "Django",
    "Git", "GitHub", "Docker", "AWS", "Azure", "Excel",
    "Pandas", "NumPy", "scikit-learn", "TensorFlow", "PyTorch",
    "Machine Learning", "Deep Learning", "NLP",
    "Data Analysis", "Data Analytics", "Power BI", "Tableau",
    "Statistics", "ETL", "dbt", "Hadoop", "Spark", "YOLO",
    "BERT", "REST API", "REST APIs", "Linux", "Agile"
]


def extract_text(file_path):
    """Read text from a PDF or DOCX resume."""
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    if path.suffix.lower() == ".pdf":
        with fitz.open(path) as pdf:
            text = "\n".join(page.get_text() for page in pdf)

    elif path.suffix.lower() == ".docx":
        document = Document(path)
        parts = [p.text for p in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        text = "\n".join(parts)

    else:
        raise ValueError("Please provide a PDF or DOCX file.")

    if not text.strip():
        raise ValueError(
            "No readable text found. The PDF may be scanned or image-based."
        )

    return text


def extract_email(text):
    match = re.search(
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        text
    )
    return match.group(0) if match else None


def extract_phone(text):
    pattern = r"(?<!\w)(?:\+?\d{1,3}[\s.-]?)?\d{10}(?!\w)"

    match = re.search(pattern, text)
    if match:
        return re.sub(r"[\s.-]", "", match.group(0))

    # Fallback for common international phone formats.
    pattern = r"(?<!\w)\+?\d[\d\s().-]{8,}\d(?!\w)"

    for match in re.finditer(pattern, text):
        digits = re.sub(r"\D", "", match.group(0))
        if 10 <= len(digits) <= 13:
            return ("+" if match.group(0).strip().startswith("+") else "") + digits

    return None


def extract_name(text, file_path=None):
    """Find a plausible name near the beginning of the resume."""
    ignored = {
        "resume", "curriculum vitae", "cv", "profile",
        "contact", "contact details", "professional summary"
    }

    for line in text.splitlines()[:12]:
        line = line.strip().strip("-|: ")

        if not line or line.lower() in ignored:
            continue
        if "@" in line or "http" in line.lower():
            continue
        if re.search(r"\d", line):
            continue
        if len(line) > 60 or len(line.split()) > 5:
            continue
        if re.search(
            r"\b(engineer|developer|analyst|student|manager|"
            r"phone|mobile|email|address|linkedin|github)\b",
            line,
            re.IGNORECASE
        ):
            continue

        words = re.findall(r"[A-Za-z][A-Za-z'-]*", line)
        if 2 <= len(words) <= 4:
            return line

    if file_path:
        return Path(file_path).stem.replace("_", " ").replace("-", " ")

    return None


def extract_skills(text):
    found = []

    for skill in SKILLS:
        pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"
        if re.search(pattern, text, re.IGNORECASE):
            found.append(skill)

    return sorted(found, key=str.lower)


def extract_link(text, platform):
    domain = (
        r"linkedin\.com"
        if platform == "linkedin"
        else r"github\.com"
    )

    pattern = rf"https?://(?:www\.)?{domain}/[^\s<>|,;]+"
    match = re.search(pattern, text, re.IGNORECASE)

    return match.group(0).rstrip(".,)") if match else None


def extract_section(text, heading_pattern):
    headings = [
        r"education|academic background|academic qualifications",
        r"work experience|professional experience|employment history|internships|experience",
        r"technical skills|skills|core competencies|technologies",
        r"projects|academic projects|personal projects",
        r"certifications|certificates",
        r"summary|professional summary|profile|objective",
        r"achievements|awards",
        r"languages|interests|references"
    ]

    lines = text.splitlines()
    capturing = False
    result = []

    for line in lines:
        current = line.strip().strip(" :-\t")
        if not current:
            continue

        is_heading = any(
            re.fullmatch(p, current, re.IGNORECASE)
            for p in headings
        )

        if is_heading:
            if capturing:
                break
            if re.fullmatch(heading_pattern, current, re.IGNORECASE):
                capturing = True
            continue

        if capturing:
            result.append(current)

    return result


def extract_education(text):
    lines = extract_section(
        text,
        r"education|academic background|academic qualifications"
    )

    results = []
    degree_pattern = (
        r"\b(B\.?\s?Tech|B\.?\s?E\.?|M\.?\s?Tech|M\.?\s?E\.?|"
        r"Bachelor(?:'s)?|Master(?:'s)?|Ph\.?\s?D\.?|MBA|"
        r"BCA|MCA|BSc|MSc|Diploma|PUC|SSLC|Higher Secondary)\b"
    )

    for line in lines:
        degree = re.search(degree_pattern, line, re.IGNORECASE)
        institution = re.search(
            r"(?:\bat\b|\bfrom\b|[-–—])\s*(.+)$",
            line,
            re.IGNORECASE
        )

        if degree or re.search(
            r"\b(university|college|institute|school)\b",
            line,
            re.IGNORECASE
        ):
            results.append({
                "degree": degree.group(0) if degree else None,
                "institution": institution.group(1).strip() if institution else None,
                "details": line
            })

    return results


def extract_experience(text):
    lines = extract_section(
        text,
        r"work experience|professional experience|employment history|internships|experience"
    )

    return [{"details": line} for line in lines if len(line) >= 4]


def extract_resume(file_path):
    text = extract_text(file_path)

    return {
        "name": extract_name(text, file_path),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "skills": extract_skills(text),
        "education": extract_education(text),
        "work_experience": extract_experience(text),
        "linkedin": extract_link(text, "linkedin"),
        "github": extract_link(text, "github")
    }